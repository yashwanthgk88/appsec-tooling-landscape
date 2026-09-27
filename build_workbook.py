import json, re, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

import os; R = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'research') + '/'
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'AppSec_Tooling_Landscape_2026.xlsx')
import datetime; AS_OF = datetime.date.today().isoformat()

CATS = [
    ('SAST', 'sast.json'),
    ('SCA', 'sca.json'),
    ('Secret Scanning', 'secrets.json'),
    ('SBOM & AI BOM', 'sbom_aibom.json'),
    ('Threat Modeling', 'threat_modeling.json'),
    ('Security Requirements', 'security_requirements.json'),
    ('Platforms & ASPM', 'platforms_aspm.json'),
]
DATA = {name: json.load(open(R + f)) for name, f in CATS}

# ---------- styles ----------
FONT = 'Arial'
NAVY = '1F2A44'; YELLOW = 'FFE600'; GREY = 'F2F2F2'; MID = 'D9D9D9'; WHITE = 'FFFFFF'
hdr_font = Font(name=FONT, bold=True, color=WHITE, size=10)
hdr_fill = PatternFill('solid', fgColor=NAVY)
body_font = Font(name=FONT, size=9)
bold = Font(name=FONT, size=9, bold=True)
title_font = Font(name=FONT, size=16, bold=True, color=NAVY)
sub_font = Font(name=FONT, size=10, italic=True, color='595959')
h2_font = Font(name=FONT, size=12, bold=True, color=NAVY)
link_font = Font(name=FONT, size=9, color='0563C1', underline='single')
input_font = Font(name=FONT, size=9, color='0000FF')
thin = Side(style='thin', color=MID)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='top', wrap_text=True)

def clean(v):
    if v is None: return ''
    if isinstance(v, list) and v and isinstance(v[0], dict) and 'date' in v[0]:
        return '\n'.join(f"{u.get('date','')}: [{u.get('type','')}] {u.get('summary','')}" for u in sorted(v, key=lambda x: x.get('date',''), reverse=True))
    if isinstance(v, bool): return 'Yes' if v else 'No'
    if isinstance(v, (int, float)): return v
    if isinstance(v, dict):
        return '\n'.join(f'{k.replace("_", " ").capitalize()}: {clean(x)}' for k, x in v.items())
    if isinstance(v, list):
        return '\n'.join('- ' + str(clean(x)) if not isinstance(x, dict) else clean(x) for x in v)
    s = str(v)
    if ' | ' in s or (s.count('|') >= 2):
        s = '\n'.join('- ' + p.strip() for p in s.split('|') if p.strip())
    return s.replace('—', ', ').replace('–', '-')

def sources_cell(v):
    s = clean(v)
    return s

def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def write_header(ws, row, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = center; cell.border = border
    ws.row_dimensions[row].height = 32

def est_height(values, widths):
    lines = 1
    for v, w in zip(values, widths):
        s = str(v) if v is not None else ''
        n = 0
        for part in s.split('\n'):
            n += max(1, math.ceil(len(part) / max(1, (w * 1.35))))
        lines = max(lines, n)
    return min(300, max(15, 12.2 * lines))

def add_ai_maturity_cf(ws, col_letter, first, last):
    rng = f'{col_letter}{first}:{col_letter}{last}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("GA",{col_letter}{first}))'], fill=PatternFill('solid', fgColor='C6EFCE')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("Beta",{col_letter}{first}))'], fill=PatternFill('solid', fgColor='FFEB9C')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("Marketing",{col_letter}{first}))'], fill=PatternFill('solid', fgColor='FFC7CE')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col_letter}{first},4)="None"'], fill=PatternFill('solid', fgColor='E7E6E6')))

# ---------- scoring rules ----------
def s_ai(m):
    m = (m or '').lower()
    if m.startswith('ga') or ' ga' in m[:12]: return 5
    if 'ga' in m and 'none' in m: return 3
    if 'beta' in m or 'preview' in m: return 3
    if 'marketing' in m: return 2
    return 1
def s_deploy(d):
    d = (d or '').lower()
    clauses = [c for c in re.split(r'[;,/()]', d) if c.strip() and not re.search(r'\bno\b|\bnot\b|legacy|limited|end of life|deprecated|unverified|only', c)]
    d = ' '.join(clauses)
    if 'air' in d: return 5
    if 'on-prem' in d or 'on prem' in d or 'self-host' in d or 'self host' in d or 'self-managed' in d or 'desktop' in d or 'cli' in d or 'local' in d: return 4
    if 'hybrid' in d or 'single-tenant' in d or 'dedicated' in d: return 3
    if 'saas' in d or 'cloud' in d: return 2
    return 1
def s_analyst(a):
    a = (a or '').lower()
    if 'not rated' in a and 'leader' not in a: return 1
    if 'leader' in a: return 5
    if 'visionar' in a or 'challenger' in a or 'strong performer' in a: return 4
    if 'niche' in a or 'contender' in a or 'peer insights' in a or 'cool vendor' in a or 'radar' in a: return 3
    if a.strip(): return 2
    return 1
def s_price(p):
    p = (p or '').lower()
    return {'high': 5, 'medium': 3, 'low': 2}.get(p.split()[0] if p.split() else '', 1)
def s_oss(t):
    t = (t or '').lower().strip()
    if t.startswith('open source') or t.startswith('framework') or t.startswith('standard') or t.startswith('free'): return 5
    if t.startswith('open-core') or t.startswith('open core'): return 3
    if 'free' in t or 'open source' in t or 'open-source' in t or 'oss' in t: return 2
    return 1

# ---------- workbook ----------
wb = Workbook()

# ===== README =====
ws = wb.active; ws.title = 'README'
ws['A1'] = 'AppSec Tooling Landscape 2026'; ws['A1'].font = title_font
ws['A2'] = f'Internal analysis for the AppSec team. Research cut-off {AS_OF}. Prepared for the 2027 priority "Better Tooling and AI Leverage".'; ws['A2'].font = sub_font
rows = [
    ('Purpose', 'Compare every credible tool per category (commercial, open-core, open source, OWASP) with a hard look at AI capabilities, so the team can decide re-tooling, AI adoption across solutions, and the 40% effort-reduction benchmark.'),
    ('Scope', 'SAST, SCA, Secret Scanning, SBOM and AI BOM, Threat Modeling (incl. AI-system threat modeling), Security Requirements, and a Platforms and ASPM tab because most vendors now bundle. 250 tool rows in total.'),
    ('How to read a category tab', 'One row per tool. Columns: identity, deployment, coverage, key features, AI capabilities (broken out by capability type), AI maturity flag, integrations, sectors and reference customers, analyst position, cons, pricing model and indicative cost, best fit, last verified, sources. Filters are on; freeze panes keep tool names visible.'),
    ('AI maturity flag', 'GA = generally available and documented by vendor. Beta or Preview = announced, gated or limited. Marketing claim = named on the website but no documentation or product behaviour found. None = no AI capability for that category. Cells are colour-coded.'),
    ('Licensing cost column', 'Public information only. No internal quotes are used. Pricing confidence: High = vendor-published list price; Medium = credible third party (marketplace listing, G2, Vendr); Low = anecdotal or single source; None = not publicly disclosed. Enterprise contracts for Checkmarx, Veracode, Black Duck, Fortify, HCL, Sonatype, JFrog are quote-only.'),
    ('Sectors and reference customers', 'Vendor-claimed. Case studies show who bought, not who is satisfied. Treat as an indicator of market presence only.'),
    ('Scoring Matrix', 'Rule-based scores derived from the data columns (see rules on that sheet), weighted by editable weights in the blue cells. Add your own adjustment in the "Team adjustment" column after hands-on evaluation. Totals are formulas and recalculate.'),
    ('OWASP', 'Every category tab is paired with the current OWASP projects and standards on the "OWASP Projects" sheet: status, latest version, link, and how the team should use it.'),
    ('New entrants', 'Tools launched or funded in 2025 to 2026 are flagged in each tab and consolidated on the "New Entrants" sheet with a reference link each.'),
    ('Verification', 'Every row carries a last-verified date and 3+ source URLs. Items the research could not confirm are marked "unverified" or "not publicly disclosed" in the cell rather than guessed. Known caveats are listed on "Sources & Caveats". Re-verify AI claims before external use; this market changes monthly.'),
    ('Legend', 'Blue text = editable input (Scoring Matrix weights and adjustments). Green fill = GA AI. Amber = Beta or Preview. Red = Marketing claim. Grey = None.'),
]
r = 4
for k, v in rows:
    ws.cell(row=r, column=1, value=k).font = bold
    c = ws.cell(row=r, column=2, value=v); c.font = body_font; c.alignment = wrap
    ws.row_dimensions[r].height = est_height([v], [110])
    r += 1
ws.cell(row=r + 1, column=1, value='Sheet index').font = h2_font
r += 2
index = [('Executive Summary', 'Top 3 per category, best OSS pick, consolidation view'),
         ('Scoring Matrix', 'All tools scored, weighted, sortable'),
         ('AI Capabilities', 'Cross-category view of every tool\'s AI capability and maturity'),
         ('SAST', 'Static analysis'), ('SCA', 'Software composition analysis'), ('Secret Scanning', 'Secrets and NHI detection'),
         ('SBOM & AI BOM', 'SBOM generation, management, AI BOM and model scanning'), ('Threat Modeling', 'Tools, AI assistants and AI-system frameworks'),
         ('Security Requirements', 'Requirements generation, standards and tooling'), ('Platforms & ASPM', 'Bundled platforms, ASPM, AI AppSec agents'),
         ('Feature Updates', 'Dated feature launches, GA announcements, acquisitions and pricing changes per tool'), ('OWASP Projects', 'Current OWASP projects per category'), ('New Entrants', '2025 to 2026 launches and funded startups, with links'),
         ('Market Notes', 'Trends, recommendations and consolidation analysis per category'), ('Sources & Caveats', 'Research method and unverified items')]
for k, v in index:
    ws.cell(row=r, column=1, value=k).font = bold
    ws.cell(row=r, column=2, value=v).font = body_font
    r += 1
set_widths(ws, [28, 120])

# ===== Category tabs =====
COLS = [
    ('#', 4), ('Tool', 26), ('Vendor', 18), ('Product URL', 30), ('New entrant (2025-26)', 10), ('Type', 12), ('Deployment', 16),
    ('Language / ecosystem coverage', 34), ('Key features', 48), ('AI capabilities (by type)', 60), ('AI maturity', 14), ('AI verified source', 30),
    ('Integrations', 34), ('Sectors and reference customers (vendor-claimed)', 36), ('Analyst position', 30), ('Cons and known gaps', 44),
    ('Pricing model', 20), ('Indicative cost (USD, public)', 26), ('Pricing confidence', 11), ('Best fit for', 34), ('Recent updates (dated)', 44), ('Last verified', 11), ('Sources', 40),
]
KEYS = ['tool', 'vendor', 'product_url', 'is_new_entrant', 'type', 'deployment', 'languages_coverage', 'key_features', 'ai_capabilities', 'ai_maturity',
        'ai_verified_source', 'integrations', 'sectors_reference_customers', 'analyst_position', 'cons_known_gaps', 'pricing_model', 'indicative_cost_usd',
        'pricing_confidence', 'best_fit_for', 'recent_updates', 'last_verified', 'sources']
score_rows = []  # for scoring matrix
ai_rows = []
new_rows = []
for cat, _ in CATS:
    d = DATA[cat]
    ws = wb.create_sheet(cat)
    ws['A1'] = f'{cat}: tool landscape'; ws['A1'].font = title_font
    ws['A2'] = f'{len(d["tools"])} tools. Research cut-off {AS_OF}. Vendor-claimed items are labelled; unverified items say so. See OWASP Projects and Market Notes sheets for this category.'; ws['A2'].font = sub_font
    write_header(ws, 4, [c[0] for c in COLS])
    widths = [c[1] for c in COLS]
    set_widths(ws, widths)
    tools = sorted(d['tools'], key=lambda t: (t.get('type', '').lower().startswith('open'), t['tool'].lower()))
    r = 5
    for i, t in enumerate(tools, 1):
        vals = [i] + [clean(t.get(k)) for k in KEYS]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = body_font; cell.alignment = wrap; cell.border = border
        ws.cell(row=r, column=2).font = bold
        url = t.get('product_url') or ''
        if url.startswith('http'):
            ws.cell(row=r, column=4).hyperlink = url; ws.cell(row=r, column=4).font = link_font
        src = t.get('ai_verified_source') or ''
        if isinstance(src, str) and src.startswith('http') and ' ' not in src.strip():
            ws.cell(row=r, column=12).hyperlink = src.strip(); ws.cell(row=r, column=12).font = link_font
        ws.cell(row=r, column=5).alignment = center; ws.cell(row=r, column=11).alignment = center; ws.cell(row=r, column=19).alignment = center; ws.cell(row=r, column=22).alignment = center
        ws.row_dimensions[r].height = est_height(vals, widths)
        score_rows.append((cat, t))
        ai_rows.append((cat, t))
        if t.get('is_new_entrant') is True:
            new_rows.append((cat, t))
        r += 1
    last = r - 1
    ws.freeze_panes = 'C5'
    ws.auto_filter.ref = f'A4:{get_column_letter(len(COLS))}{last}'
    add_ai_maturity_cf(ws, 'K', 5, last)
    ws.conditional_formatting.add(f'E5:E{last}', CellIsRule(operator='equal', formula=['"Yes"'], fill=PatternFill('solid', fgColor=YELLOW)))
    ws.sheet_view.zoomScale = 90

# ===== AI Capabilities =====
ws = wb.create_sheet('AI Capabilities', 3)
ws['A1'] = 'AI capabilities across all categories'; ws['A1'].font = title_font
ws['A2'] = 'One row per tool per category. Sort or filter by maturity. Capability text is grouped by type (triage, autofix, rules, reachability, NL query, agent/MCP, AI-code scanning, model used).'; ws['A2'].font = sub_font
hdrs = ['Category', 'Tool', 'Vendor', 'Type', 'AI maturity', 'AI capabilities (by type)', 'AI verified source', 'Last verified']
write_header(ws, 4, hdrs); w = [18, 28, 18, 12, 14, 90, 36, 11]; set_widths(ws, w)
r = 5
for cat, t in ai_rows:
    vals = [cat, t['tool'], t.get('vendor'), t.get('type'), clean(t.get('ai_maturity')), clean(t.get('ai_capabilities')), clean(t.get('ai_verified_source')), t.get('last_verified')]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v); cell.font = body_font; cell.alignment = wrap; cell.border = border
    ws.cell(row=r, column=2).font = bold
    src = t.get('ai_verified_source') or ''
    if isinstance(src, str) and src.startswith('http') and ' ' not in src.strip():
        ws.cell(row=r, column=7).hyperlink = src.strip(); ws.cell(row=r, column=7).font = link_font
    ws.row_dimensions[r].height = est_height(vals, w)
    r += 1
ws.freeze_panes = 'C5'; ws.auto_filter.ref = f'A4:H{r-1}'; add_ai_maturity_cf(ws, 'E', 5, r - 1)
ai_counts = {}
for cat, t in ai_rows:
    m = clean(t.get('ai_maturity')); k = 'GA' if m.startswith('GA') else 'Beta or Preview' if ('Beta' in m or 'Preview' in m) else 'Marketing claim' if 'Marketing' in m else 'None' if m.startswith('None') else 'Mixed'
    ai_counts.setdefault(cat, {}).setdefault(k, 0); ai_counts[cat][k] += 1

# ===== Scoring Matrix =====
ws = wb.create_sheet('Scoring Matrix', 2)
ws['A1'] = 'Scoring matrix (rule-based, weighted, editable)'; ws['A1'].font = title_font
ws['A2'] = 'Scores 1 to 5 are derived from the data columns by the rules below. Change the blue weights or add a Team adjustment; totals recalculate. Sort by Weighted score after filtering to a category.'; ws['A2'].font = sub_font
ws['A4'] = 'Weights (edit the blue cells; they should sum to 100%)'; ws['A4'].font = h2_font
crit = [('AI maturity', 0.30, 'GA = 5; Beta/Preview or GA-for-platform-but-not-this-category = 3; Marketing claim = 2; None = 1'),
        ('Deployment flexibility', 0.20, 'Air-gap supported = 5; on-prem, self-hosted, self-managed, desktop, local or CLI = 4; hybrid, single-tenant or dedicated = 3; SaaS only = 2; N/A = 1. Negated clauses ("no on-prem", "legacy", "limited") are ignored'),
        ('Analyst standing', 0.20, 'Leader in a Gartner MQ or Forrester Wave = 5; Visionary, Challenger or Strong Performer = 4; Niche, Contender, Cool Vendor, GigaOm or Peer Insights only = 3; other mention = 2; not rated = 1'),
        ('Pricing transparency', 0.15, 'Pricing confidence High = 5; Medium = 3; Low = 2; None = 1'),
        ('Open source / free availability', 0.15, 'Open source, framework, standard or free = 5; open-core = 3; commercial with a free tier or open-source component = 2; commercial only = 1')]
write_header(ws, 5, ['Criterion', 'Weight', 'Scoring rule'])
for i, (n, wgt, rule) in enumerate(crit):
    ws.cell(row=6 + i, column=1, value=n).font = bold
    c = ws.cell(row=6 + i, column=2, value=wgt); c.font = input_font; c.number_format = '0%'; c.fill = PatternFill('solid', fgColor='FFF2CC')
    c = ws.cell(row=6 + i, column=3, value=rule); c.font = body_font; c.alignment = wrap
    ws.row_dimensions[6 + i].height = 28
ws.cell(row=11, column=1, value='Total').font = bold
ws.cell(row=11, column=2, value='=SUM(B6:B10)').number_format = '0%'; ws.cell(row=11, column=2).font = bold
ws.cell(row=12, column=1, value='Note').font = bold
ws.cell(row=12, column=3, value='Rule-based scores measure evidence in the public record, not tool quality. A niche OSS tool can outscore a Leader on transparency and deployment. Use "Team adjustment" (-2 to +2) after a hands-on bake-off, e.g. OWASP Benchmark results for SAST.').font = body_font
ws.cell(row=12, column=3).alignment = wrap; ws.row_dimensions[12].height = 40
HR = 14
hdrs = ['Category', 'Tool', 'Vendor', 'Type', 'AI maturity (as recorded)', 'AI maturity score', 'Deployment score', 'Analyst score', 'Pricing transparency score', 'OSS / free score', 'Team adjustment (-2 to +2)', 'Weighted score (1-5)', 'Rank in category']
write_header(ws, HR, hdrs); set_widths(ws, [18, 30, 18, 12, 22, 10, 10, 10, 10, 10, 12, 12, 10])
r = HR + 1
first = r
for cat, t in score_rows:
    vals = [cat, t['tool'], t.get('vendor'), t.get('type'), clean(t.get('ai_maturity')), s_ai(clean(t.get('ai_maturity'))), s_deploy(t.get('deployment')), s_analyst(t.get('analyst_position')), s_price(t.get('pricing_confidence')), s_oss(t.get('type'))]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v); cell.font = body_font; cell.alignment = wrap if c <= 5 else center; cell.border = border
    ws.cell(row=r, column=2).font = bold
    c = ws.cell(row=r, column=11, value=0); c.font = input_font; c.alignment = center; c.border = border; c.fill = PatternFill('solid', fgColor='FFF2CC')
    c = ws.cell(row=r, column=12, value=f'=ROUND(SUMPRODUCT(F{r}:J{r},TRANSPOSE($B$6:$B$10))/$B$11+K{r},2)')
    c.font = bold; c.alignment = center; c.border = border; c.number_format = '0.00'
    r += 1
last = r - 1
for rr in range(first, last + 1):
    c = ws.cell(row=rr, column=13, value=f'=COUNTIFS($A${first}:$A${last},A{rr},$L${first}:$L${last},">"&L{rr})+1')
    c.font = body_font; c.alignment = center; c.border = border
ws.freeze_panes = f'C{HR+1}'; ws.auto_filter.ref = f'A{HR}:M{last}'
ws.conditional_formatting.add(f'L{first}:L{last}', FormulaRule(formula=[f'L{first}>=4'], fill=PatternFill('solid', fgColor='C6EFCE')))
ws.conditional_formatting.add(f'L{first}:L{last}', FormulaRule(formula=[f'AND(L{first}>=3,L{first}<4)'], fill=PatternFill('solid', fgColor='FFEB9C')))
add_ai_maturity_cf(ws, 'E', first, last)

# ===== Executive Summary =====
ws = wb.create_sheet('Executive Summary', 1)
ws['A1'] = 'Executive summary'; ws['A1'].font = title_font
ws['A2'] = f'Research cut-off {AS_OF}. Recommendations are for an enterprise consulting delivery team; they are starting points for a bake-off, not a procurement decision.'; ws['A2'].font = sub_font
r = 4
ws.cell(row=r, column=1, value='Coverage and AI maturity by category').font = h2_font; r += 1
write_header(ws, r, ['Category', 'Tools analysed', 'AI: GA', 'AI: Beta or Preview', 'AI: Marketing claim', 'AI: None', 'AI: Mixed / other', 'New entrants 2025-26', 'OSS / open-core rows']); r += 1
tstart = r
for cat, _ in CATS:
    tl = DATA[cat]['tools']; ac = ai_counts.get(cat, {})
    vals = [cat, len(tl), ac.get('GA', 0), ac.get('Beta or Preview', 0), ac.get('Marketing claim', 0), ac.get('None', 0), ac.get('Mixed', 0),
            sum(1 for t in tl if t.get('is_new_entrant') is True), sum(1 for t in tl if s_oss(t.get('type')) >= 3)]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v); cell.font = body_font; cell.border = border; cell.alignment = center if c > 1 else wrap
    ws.cell(row=r, column=1).font = bold
    r += 1
ws.cell(row=r, column=1, value='Total').font = bold
for c in range(2, 10):
    cell = ws.cell(row=r, column=c, value=f'=SUM({get_column_letter(c)}{tstart}:{get_column_letter(c)}{r-1})'); cell.font = bold; cell.alignment = center; cell.border = border
r += 2
ws.cell(row=r, column=1, value='Top contenders and best open-source pick per category').font = h2_font; r += 1
write_header(ws, r, ['Category', 'Rank', 'Tool', 'Reasoning', 'Best OSS / free pick']); r += 1
def find_top3(notes):
    for k, v in notes.items():
        if 'top' in k.lower() and isinstance(v, list) and v and isinstance(v[0], dict): return v
    return []
def find_oss(notes):
    for k, v in notes.items():
        if 'oss' in k.lower() or 'free' in k.lower(): return clean(v)
    return ''
for cat, _ in CATS:
    notes = DATA[cat]['category_notes']
    tops = find_top3(notes); oss = find_oss(notes)
    if cat == 'SBOM & AI BOM':
        tops = [dict(x, tool='[SBOM] ' + x.get('tool', '')) for x in notes.get('top3_sbom_contenders', [])] + [dict(x, tool='[AI BOM] ' + x.get('tool', '')) for x in notes.get('top3_aibom_contenders', [])]
    start = r
    tops = [x for x in tops if (x.get('tool') or x.get('name'))] + [dict(rank='', tool='Also considered', reasoning=clean(x.get('honourable_mentions') or x.get('alternates') or x.get('alternate') or x)) for x in tops if not (x.get('tool') or x.get('name'))]
    oss_lines = max(1, len(oss) / 45)
    for i, x in enumerate(tops):
        vals = [cat if i == 0 else '', x.get('rank', i + 1), x.get('tool') or x.get('name'), clean(x.get('reasoning') or x.get('reason') or x.get('why')), oss if i == 0 else '']
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v); cell.font = body_font; cell.alignment = wrap; cell.border = border
        ws.cell(row=r, column=1).font = bold; ws.cell(row=r, column=3).font = bold; ws.cell(row=r, column=2).alignment = center
        ws.row_dimensions[r].height = max(est_height([vals[3]], [90]), 12.5 * oss_lines / max(1, len(tops)) + 4)
        r += 1
    if len(tops) > 1:
        ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
        ws.merge_cells(start_row=start, start_column=5, end_row=r - 1, end_column=5)
r += 1
ws.cell(row=r, column=1, value='Consolidate vs best-of-breed (from Platforms & ASPM research)').font = h2_font; r += 1
cvb = DATA['Platforms & ASPM']['category_notes'].get('consolidate_vs_best_of_breed', {})
for k, v in (cvb.items() if isinstance(cvb, dict) else [('analysis', cvb)]):
    ws.cell(row=r, column=1, value=k.replace('_', ' ').capitalize()).font = bold
    c = ws.cell(row=r, column=2, value=clean(v)); c.font = body_font; c.alignment = wrap
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.row_dimensions[r].height = est_height([clean(v)], [150])
    r += 1
r += 1
ws.cell(row=r, column=1, value='Reading the 2027 priorities against this data').font = h2_font; r += 1
obs = [
    ('Re-tool to best in class', 'Checkmarx One is the only vendor verified as Leader in Gartner MQ AST (Oct 2025), Gartner MQ Software Supply Chain Security (Jun 2026) and the Forrester SAST Wave (Q3 2025). Black Duck and Sonatype lead on SCA and SBOM depth for regulated and air-gapped clients. GitHub and GitLab win where the client already owns the SCM. Snyk leads on developer experience and the agentic roadmap (Evo).'),
    ('AI across all solutions', 'AI autofix and AI triage are now GA at every MQ Leader. MCP servers exist from 15+ vendors. The gap is in secrets (incumbents have no secrets-specific AI; GitGuardian, Cycode and GitHub do), threat modeling (only ThreatModeler Nexus, IriusRisk Jeff, Devici, Threat Canvas and STRIDE GPT generate models from code or IaC), and AI BOM (only cdxgen, Snyk aibom, OWASP AIBOM Generator, ZeroPath and agent-bom emit standards-based CycloneDX AI BOMs).'),
    ('40% effort reduction', 'Vendor-reported figures only: Semgrep Assistant 96% triage agreement, Endor Labs 95% false-positive cut, Corgea 90% noise reduction, Betterleaks 98.6% recall. None are independently benchmarked. Run OWASP Benchmark (SAST) and a fixed corpus of internal repos through the shortlist to produce the team\'s own baseline before quoting a number.'),
    ('Where the incumbents are weak', 'Veracode, Fortify, HCL and Black Duck have no dedicated secrets product, no validity checking, no NHI inventory. Fortify Aviator and HCL RapidFix are SAST-only. Veracode SBOM API still emits CycloneDX 1.4. Fortify SBOM export is Enterprise-tier only. Pricing is quote-only across all of them.'),
    ('OWASP baseline', 'ASVS 5.0 (May 2025), OWASP Top 10:2025 (new A03 Software Supply Chain Failures), CycloneDX 1.7 (Oct 2025, ECMA-424), Dependency-Track 5.x (Jun 2026), LLM Top 10 2026, Agentic Top 10 2026 and the Agent Control Standard. Dependency-Check and SecurityRAT are effectively legacy. Threat Dragon has no AI roadmap.'),
]
for k, v in obs:
    ws.cell(row=r, column=1, value=k).font = bold
    c = ws.cell(row=r, column=2, value=v); c.font = body_font; c.alignment = wrap
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.row_dimensions[r].height = est_height([v], [150])
    r += 1
set_widths(ws, [26, 12, 34, 95, 55])

# ===== Feature Updates =====
ws = wb.create_sheet('Feature Updates')
ws['A1'] = 'Feature updates, launches and vendor news per tool'; ws['A1'].font = title_font
ws['A2'] = 'One row per dated update recorded against a tool (feature launches, GA announcements, acquisitions, pricing changes, releases). Newest first. Populated by the weekly update run.'; ws['A2'].font = sub_font
write_header(ws, 4, ['Date', 'Category', 'Tool', 'Update type', 'What changed', 'Source'])
w = [11, 18, 30, 16, 90, 40]; set_widths(ws, w)
r = 5
upd = []
for cat, _ in CATS:
    for t in DATA[cat]['tools']:
        for u in (t.get('recent_updates') or []):
            if isinstance(u, dict):
                upd.append([u.get('date', ''), cat, t['tool'], u.get('type', ''), u.get('summary', ''), u.get('source', '')])
upd.sort(key=lambda x: x[0], reverse=True)
for vals in upd:
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v); cell.font = body_font; cell.alignment = wrap; cell.border = border
    ws.cell(row=r, column=3).font = bold
    if isinstance(vals[5], str) and vals[5].startswith('http'):
        ws.cell(row=r, column=6).hyperlink = vals[5]; ws.cell(row=r, column=6).font = link_font
    ws.row_dimensions[r].height = est_height(vals, w); r += 1
if not upd:
    ws.cell(row=5, column=1, value='No dated updates recorded yet. The weekly run appends entries here.').font = body_font
ws.freeze_panes = 'A5'; ws.auto_filter.ref = f'A4:F{max(r-1,5)}'

# ===== OWASP Projects =====
ws = wb.create_sheet('OWASP Projects')
ws['A1'] = 'OWASP projects and standards per category'; ws['A1'].font = title_font
ws['A2'] = 'Current status, latest version and how the team should use each. Verified against owasp.org and project GitHub repositories where reachable.'; ws['A2'].font = sub_font
write_header(ws, 4, ['Category', 'Project', 'Status', 'Latest version / date', 'Link', 'Details', 'How the team should use it'])
w = [18, 30, 24, 34, 36, 70, 60]; set_widths(ws, w)
r = 5
def owasp_rows(cat, ow):
    out = []
    if isinstance(ow, dict) and 'projects' in ow and isinstance(ow['projects'], list):
        items = [(p.get('project') or p.get('name') or '', p) for p in ow['projects']]
        extra = {k: v for k, v in ow.items() if k != 'projects'}
        if extra: items.append(('Category summary', extra))
    elif isinstance(ow, dict):
        items = []
        for k, v in ow.items():
            if isinstance(v, dict): items.append((v.get('name') or v.get('project') or k.replace('_', ' ').upper() if len(k) < 6 else v.get('name') or k.replace('_', ' ').title(), v))
            elif isinstance(v, list):
                for p in v: items.append((p.get('project') or p.get('name') or k, p) if isinstance(p, dict) else (k, {'details': p}))
            else: items.append((k.replace('_', ' ').title(), {'details': v}))
    else:
        items = [((p.get('project') or p.get('name') or ''), p) for p in ow]
    for name, p in items:
        if not isinstance(p, dict): p = {'details': p}
        status = p.get('status', '')
        ver = p.get('latest_version') or p.get('latest_release') or p.get('latest_version_date') or p.get('version') or ''
        link = p.get('link') or p.get('url') or ''
        how = p.get('how_to_use') or p.get('how_team_should_use') or p.get('how_the_team_should_use') or p.get('usage') or p.get('use') or ''
        skip = {'status', 'latest_version', 'latest_release', 'latest_version_date', 'version', 'link', 'url', 'how_to_use', 'how_team_should_use', 'how_the_team_should_use', 'usage', 'use', 'name', 'project'}
        details = '\n'.join(f'{k.replace("_", " ").capitalize()}: {clean(v)}' for k, v in p.items() if k not in skip)
        out.append([cat, name, clean(status), clean(ver), link, details, clean(how)])
    return out
for cat, _ in CATS:
    for vals in owasp_rows(cat, DATA[cat]['owasp']):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v); cell.font = body_font; cell.alignment = wrap; cell.border = border
        ws.cell(row=r, column=2).font = bold
        if isinstance(vals[4], str) and vals[4].startswith('http'):
            ws.cell(row=r, column=5).hyperlink = vals[4]; ws.cell(row=r, column=5).font = link_font
        ws.row_dimensions[r].height = est_height(vals, w)
        r += 1
ws.freeze_panes = 'C5'; ws.auto_filter.ref = f'A4:G{r-1}'

# ===== New Entrants =====
ws = wb.create_sheet('New Entrants')
ws['A1'] = 'New entrants 2025 to 2026'; ws['A1'].font = title_font
ws['A2'] = 'Tools launched, forked or first funded in 2025-2026, flagged in the category tabs, plus additional names surfaced in each sweep that did not merit a full row. Funding figures are from press coverage and are marked where unverified.'; ws['A2'].font = sub_font
write_header(ws, 4, ['Category', 'Tool', 'Vendor', 'Reference link', 'Why it is on the list', 'AI maturity', 'Full row in category tab?'])
w = [18, 28, 20, 40, 80, 14, 12]; set_widths(ws, w)
r = 5
seen = set()
for cat, t in new_rows:
    why = clean(t.get('is_new_entrant_basis') or t.get('best_fit_for') or '')
    vals = [cat, t['tool'], t.get('vendor'), t.get('product_url'), why, clean(t.get('ai_maturity')), 'Yes']
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v); cell.font = body_font; cell.alignment = wrap; cell.border = border
    ws.cell(row=r, column=2).font = bold
    if isinstance(vals[3], str) and vals[3].startswith('http'):
        ws.cell(row=r, column=4).hyperlink = vals[3]; ws.cell(row=r, column=4).font = link_font
    ws.row_dimensions[r].height = est_height(vals, w)
    seen.add((cat, t['tool'].lower()[:12])); r += 1
for cat, _ in CATS:
    notes = DATA[cat]['category_notes']
    for k, v in notes.items():
        if 'new_entrant' in k and isinstance(v, list):
            for x in v:
                if isinstance(x, dict):
                    name = x.get('tool') or x.get('name') or ''
                    if (cat, name.lower()[:12]) in seen: continue
                    vals = [cat, name, x.get('vendor', ''), x.get('link') or x.get('url') or '', clean(x.get('reason') or x.get('note') or x.get('status') or ''), '', 'No (sweep note only)']
                else:
                    vals = [cat, str(x).split('(')[0].strip()[:60], '', '', clean(x), '', 'No (sweep note only)']
                for c, vv in enumerate(vals, 1):
                    cell = ws.cell(row=r, column=c, value=vv); cell.font = body_font; cell.alignment = wrap; cell.border = border
                ws.cell(row=r, column=2).font = bold
                if isinstance(vals[3], str) and vals[3].startswith('http'):
                    ws.cell(row=r, column=4).hyperlink = vals[3]; ws.cell(row=r, column=4).font = link_font
                ws.row_dimensions[r].height = est_height(vals, w)
                r += 1
ws.freeze_panes = 'C5'; ws.auto_filter.ref = f'A4:G{r-1}'; add_ai_maturity_cf(ws, 'F', 5, r - 1)

# ===== Market Notes =====
ws = wb.create_sheet('Market Notes')
ws['A1'] = 'Market notes per category'; ws['A1'].font = title_font
ws['A2'] = 'Trends 2025-2026, recommendations and reasoning, best OSS picks, and analysis notes as recorded in each research stream.'; ws['A2'].font = sub_font
write_header(ws, 4, ['Category', 'Topic', 'Content'])
w = [18, 30, 150]; set_widths(ws, w)
r = 5
for cat, _ in CATS:
    for k, v in DATA[cat]['category_notes'].items():
        if k in ('last_verified',): continue
        vals = [cat, k.replace('_', ' ').capitalize(), clean(v)]
        for c, vv in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=vv); cell.font = body_font; cell.alignment = wrap; cell.border = border
        ws.cell(row=r, column=2).font = bold
        ws.row_dimensions[r].height = est_height(vals, w)
        r += 1
ws.freeze_panes = 'C5'; ws.auto_filter.ref = f'A4:C{r-1}'

# ===== Sources & Caveats =====
ws = wb.create_sheet('Sources & Caveats')
ws['A1'] = 'Research method, sources and caveats'; ws['A1'].font = title_font
r = 3
method = [
    ('Method', 'Seven parallel research streams (one per category), each run against vendor documentation, release notes, press releases, analyst press coverage (Gartner MQ AST Oct 2025, Gartner MQ Software Supply Chain Security Jun 2026, Forrester SAST Wave Q3 2025, Forrester SCA Wave Q4 2024), OWASP project pages and GitHub repositories, plus a dedicated 2025-2026 new-entrant sweep per category. Incumbents (Checkmarx, Veracode, Fortify, HCL, Black Duck, Snyk, GitHub, GitLab, Sonatype, JFrog, Mend) were back-filled on every tab they play in, including where they have no dedicated product, so AI capabilities can be compared like for like.'),
    ('What "unverified" means', 'The research could not reach a primary source in-session (some vendor pages are JavaScript-rendered or blocked to automated fetch; the web-search budget ran out in some streams). The cell says so rather than guessing. Analyst placements below Leader tier are mostly unverified because Gartner and Forrester gate the full reports.'),
    ('Pricing', 'Only vendor-published list prices carry High confidence (GitHub Code Security $30 and Secret Protection $19 per active committer per month; Snyk Team $25 per dev per month; Semgrep $15 per contributor; Aikido $0/$300/$600 per month tiers; SonarQube Cloud Team from $34 per month; SD Elements Enterprise from $75K per year; DefectDojo Pro from $100 per month; ZeroPath $1,000 per month plus $60 per dev). Vendr medians for Checkmarx, Veracode and Semgrep are Medium. Everything else is quote-only.'),
    ('AI accuracy claims', 'All accuracy or noise-reduction percentages are vendor-reported and not independently benchmarked. Treat them as marketing until reproduced on the team\'s own corpus.'),
    ('Freshness', 'AI features in this market change monthly. Every row carries a last-verified date. Re-verify before any external use, and re-run the new-entrant sweep quarterly.'),
]
for k, v in method:
    ws.cell(row=r, column=1, value=k).font = bold
    c = ws.cell(row=r, column=2, value=v); c.font = body_font; c.alignment = wrap
    ws.row_dimensions[r].height = est_height([v], [150]); r += 1
r += 1
ws.cell(row=r, column=1, value='Caveats recorded per research stream').font = h2_font; r += 1
write_header(ws, r, ['Category', 'Caveat / unverified item']); r += 1
for cat, _ in CATS:
    notes = DATA[cat]['category_notes']
    for k, v in notes.items():
        if any(x in k for x in ('caveat', 'unverified', 'verification', 'excluded', 'coverage_note', 'sweep_notes')):
            items = v if isinstance(v, list) else [v]
            for it in items:
                ws.cell(row=r, column=1, value=cat).font = bold
                c = ws.cell(row=r, column=2, value=clean(it)); c.font = body_font; c.alignment = wrap
                ws.row_dimensions[r].height = est_height([clean(it)], [150]); r += 1
set_widths(ws, [24, 150])

# order sheets
order = ['README', 'Executive Summary', 'Scoring Matrix', 'AI Capabilities'] + [c for c, _ in CATS] + ['Feature Updates', 'OWASP Projects', 'New Entrants', 'Market Notes', 'Sources & Caveats']
wb._sheets = [wb[n] for n in order]
for s in wb.worksheets:
    s.sheet_properties.tabColor = YELLOW if s.title in ('README', 'Executive Summary', 'Scoring Matrix', 'AI Capabilities') else NAVY if s.title in [c for c, _ in CATS] else '7F7F7F'
wb.save(OUT)
print('saved', OUT, 'rows', len(score_rows), 'new', len(new_rows))
